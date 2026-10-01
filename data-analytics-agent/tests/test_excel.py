"""Workbook choices survive review and the existing analytical workflow."""

from dataclasses import replace
from datetime import datetime
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from fastapi.testclient import TestClient
from openpyxl import Workbook
import pyarrow.parquet as pq
import pytest

from data_analytics_agent.api import Services, create_app
from data_analytics_agent.excel import ExcelSelection, inspect_workbook
from data_analytics_agent.uploads import (
    stage_upload,
    confirm_upload,
    UploadReview,
    get_upload,
)


def workbook_bytes(*, formula=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Sales"
    sheet.append(["Report title"])
    sheet.append(["id", "date", "amount"])
    sheet.append([7, datetime(2026, 9, 1), "=2+2" if formula else 4])
    sheet["A3"].number_format = "00000"
    sheet.append(["00008", datetime(2026, 9, 2), 6])
    sheet.append(["Total", None, 10])
    book.create_sheet("Notes").append(["not data"])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_explicit_excel_selection_review_reopen_and_sql(test_settings):
    from types import SimpleNamespace
    from data_analytics_agent.agents.text_to_sql.tools import (
        create_query_saved_results_tool,
    )

    s = Services(settings=test_settings)
    content = workbook_bytes()
    assert [i["name"] for i in inspect_workbook(content, s.settings)["sheets"]] == [
        "Sales",
        "Notes",
    ]
    selection = ExcelSelection(sheet="Sales", cell_range="A2:C4")
    with TestClient(create_app(s)) as api:
        assert (
            api.post(
                "/api/uploads", params={"filename": "sales.xlsx"}, content=content
            ).status_code
            == 422
        )
        response = api.post(
            "/api/uploads",
            params={"filename": "sales.xlsx", **selection.model_dump()},
            content=content,
        )
        assert response.status_code == 201, response.text
        u = get_upload(s.storage, response.json()["source_id"])
    confirmed = confirm_upload(
        s,
        u.thread_id,
        UploadReview(
            types={c.name: c.suggested_type for c in u.columns},
            key_columns=["id"],
            grain="One sale",
        ),
    )
    result = s.results.get(confirmed.result_id, u.thread_id)
    assert result.rows[0]["id"] == "00007"
    assert result.row_count == 2
    assert result.rows[0]["date"] == datetime(2026, 9, 1)
    assert result.upload_provenance.excel_selection == selection.model_dump()
    reopened = Services(settings=test_settings)
    assert get_upload(reopened.storage, u.source_id) == confirmed
    tool = create_query_saved_results_tool(s.results, s.runs, source_id=u.source_id)
    run = s.runs.create(u.thread_id, u.source_id, "Sum sales")
    s.runs.begin_assignment(run, "sum", "text-to-sql", "Sum sales")
    runtime = SimpleNamespace(
        state={
            "thread_id": u.thread_id,
            "run_id": run,
            "source_id": u.source_id,
            "assignment_id": "sum",
            "question": "Sum sales",
        },
        tool_call_id="sum",
    )
    output = tool.func(
        query="SELECT SUM(amount) AS total FROM sales",
        bindings={"sales": confirmed.result_id},
        purpose="Total sales",
        runtime=runtime,
    )
    assert output["row_count"] == 1
    saved = s.results.get(output["result_id"], u.thread_id)
    assert saved.rows == [{"total": 10}]
    assert pq.read_table(result.parquet_path).column("id").to_pylist() == [
        "00007",
        "00008",
    ]


def test_excel_formula_cache_requires_choice_and_value(test_settings):
    s = Services(settings=test_settings)
    content = workbook_bytes(formula=True)
    with pytest.raises(ValueError, match="Explicitly accept"):
        stage_upload(
            s, content, "sales.xlsx", ExcelSelection(sheet="Sales", cell_range="A2:C4")
        )
    with pytest.raises(ValueError, match="no stored value"):
        stage_upload(
            s,
            content,
            "sales.xlsx",
            ExcelSelection(sheet="Sales", cell_range="A2:C4", use_cached_formulas=True),
        )
    out = BytesIO()
    with ZipFile(BytesIO(content)) as src, ZipFile(out, "w", ZIP_DEFLATED) as dest:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<f>2+2</f><v />", b"<f>2+2</f><v>4</v>")
            dest.writestr(info, data)
    u = stage_upload(
        s,
        out.getvalue(),
        "sales.xlsx",
        ExcelSelection(sheet="Sales", cell_range="A2:C4", use_cached_formulas=True),
    )
    assert s.results.get(u.result_id, u.thread_id).rows[0]["amount"] == 4
    assert any("freshness is unknown" in w for w in u.warnings)


def test_excel_limits_headers_merges_and_keys(test_settings):
    s = Services(settings=replace(test_settings, max_result_rows=1))
    with pytest.raises(ValueError, match="row limit"):
        stage_upload(
            s,
            workbook_bytes(),
            "sales.xlsx",
            ExcelSelection(sheet="Sales", cell_range="A2:C4"),
        )
    s = Services(settings=test_settings)
    with pytest.raises(ValueError, match="header"):
        stage_upload(
            s,
            workbook_bytes(),
            "sales.xlsx",
            ExcelSelection(sheet="Sales", cell_range="A1:C4"),
        )
    with pytest.raises(ValueError, match="rectangular"):
        stage_upload(
            s,
            workbook_bytes(),
            "sales.xlsx",
            ExcelSelection(sheet="Sales", cell_range="A:C"),
        )
    book = Workbook()
    book.active.merge_cells("A1:B1")
    buf = BytesIO()
    book.save(buf)
    with pytest.raises(ValueError, match="merged"):
        stage_upload(
            s,
            buf.getvalue(),
            "merged.xlsx",
            ExcelSelection(sheet="Sheet", cell_range="A1:B2"),
        )
    u = stage_upload(
        s,
        workbook_bytes(),
        "sales.xlsx",
        ExcelSelection(sheet="Sales", cell_range="A2:C5"),
    )
    assert u.row_count == 3  # footer retained until user deliberately excludes it
    with pytest.raises(ValueError, match="missing values"):
        confirm_upload(
            s,
            u.thread_id,
            UploadReview(
                types={c.name: c.suggested_type for c in u.columns},
                key_columns=["date"],
            ),
        )


def test_excel_selection_ui_and_merged_title_outside_table(test_settings):
    from streamlit.testing.v1 import AppTest

    book = Workbook()
    ws = book.active
    ws.title = "Data"
    ws.merge_cells("A1:C1")
    ws["A1"] = "Sales report"
    ws.append(["id", "amount"])
    ws.append(["0001", 5])
    buf = BytesIO()
    book.save(buf)
    content = buf.getvalue()
    s = Services(settings=test_settings)
    inspect_workbook(content, s.settings)
    u = stage_upload(
        s, content, "title.xlsx", ExcelSelection(sheet="Data", cell_range="A2:B3")
    )
    assert u.row_count == 1
    assert s.results.get(u.result_id, u.thread_id).rows == [{"id": "0001", "amount": 5}]
    app = AppTest.from_string("""
import streamlit as st
from data_analytics_agent.ui.uploads import render_workbook_selection
class Client:
    def upload_file(self, filename, content, selection):
        st.session_state['submitted_selection'] = selection
        return {'thread_id':'chosen-table'}
st.session_state.setdefault('pending_workbook', {'filename':'sales.xlsx','content':b'bytes','question':'Sum sales','info':{'sheets':[{'name':'First','range':'A1:B3','preview':[['id','amount'],['001','5']]},{'name':'Second','range':'A2:C5','preview':[['Title']]}]}})
if 'submitted_selection' not in st.session_state:
    render_workbook_selection(Client())
""").run()
    assert not app.exception
    app.selectbox[0].select("Second").run()
    assert [n.value for n in app.number_input] == [2, 5]
    assert [t.value for t in app.text_input] == ["A", "C"]
    assert app.success[0].value == "Selected data: 3 rows · 3 columns"
    assert app.dataframe[0].value.index.name == "Excel row"
    assert "submitted_selection" not in app.session_state
    app.number_input[1].set_value(2).run()
    assert app.button[0].disabled
    assert "below the row containing column names" in app.error[0].value
    app.number_input[1].set_value(4)
    app.text_input[1].set_value("B").run()
    assert app.success[0].value == "Selected data: 2 rows · 2 columns"
    app.button[0].click().run()
    assert not app.exception
    assert app.session_state["submitted_selection"]["sheet"] == "Second"
    assert app.session_state["submitted_selection"]["cell_range"] == "A2:B4"
    assert app.session_state["upload_questions"]["chosen-table"] == "Sum sales"


async def test_excel_through_real_agent_harness_chart_report_and_download(
    test_settings, monkeypatch
):
    from dataclasses import replace
    from pathlib import Path
    import shutil
    from data_analytics_agent import coordinator
    from tests.test_uploads import UploadAnalyst

    monkeypatch.setattr(
        coordinator, "_build_chat_model", lambda *a, **k: UploadAnalyst()
    )
    monkeypatch.setenv("LANGSMITH_TRACING", "false")
    root = Path(__file__).parents[1]
    shutil.copy(root / "AGENTS.md", test_settings.project_root / "AGENTS.md")
    shutil.copytree(root / "skills", test_settings.project_root / "skills")
    s = Services(
        settings=replace(
            test_settings, require_sql_approval=False, require_python_approval=False
        )
    )
    book = Workbook()
    book.active.title = "Sales"
    book.active.append(["id", "amount"])
    for row in [("001", 2), ("002", 4), ("003", 6)]:
        book.active.append(row)
    buf = BytesIO()
    book.save(buf)
    u = stage_upload(
        s,
        buf.getvalue(),
        "sales.xlsx",
        ExcelSelection(sheet="Sales", cell_range="A1:B4"),
    )
    u = confirm_upload(
        s,
        u.thread_id,
        UploadReview(types={c.name: c.suggested_type for c in u.columns}),
    )

    def forbidden(*args, **kwargs):
        raise AssertionError("Excel has no warehouse access")

    monkeypatch.setattr(s, "backend_for_source", forbidden)
    run = s.runs.create(u.thread_id, u.source_id, "Sum all amounts")
    s.conversations.begin_run(u.thread_id, run)
    await s.manager().start(run)
    state = s.runs.get(run)
    assert state.status == "completed", state.error
    assert state.answer.answer == "The total is 12."
    assert (
        "sales.xlsx" in s.reports.get(state.answer.report.report_id, u.thread_id).html
    )
    with TestClient(create_app(s)) as api:
        response = api.get(
            f"/api/runs/{run}/download",
            params={"report_id": state.answer.report.report_id},
        )
        assert response.status_code == 200
        from zipfile import ZipFile

        with ZipFile(BytesIO(response.content)) as zip:
            import json

            manifest = json.loads(zip.read("manifest.json"))
            assert any(
                d.get("upload_provenance", {}).get("excel_selection", {}).get("sheet")
                == "Sales"
                for d in manifest["datasets"].values()
                if d.get("upload_provenance")
            )
