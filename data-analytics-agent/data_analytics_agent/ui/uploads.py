"""Native file entry and explicit schema/grain review before analytical work."""

import hashlib

from openpyxl.utils.cell import get_column_letter, range_boundaries
import pandas as pd
import streamlit as st

from data_analytics_agent.ui.api_client import APIError
from data_analytics_agent.uploads import MAX_UPLOAD_COLUMNS, TYPE_LABELS


def chat_submission(placeholder, *, key, max_bytes, client):
    """One native composer for attachments, editable text, and voice recording."""
    from data_analytics_agent.ui.voice import (
        prepare_voice_draft,
        review_voice_submission,
    )

    with st.bottom:
        prepare_voice_draft(client, key)
        submission = st.chat_input(
            placeholder,
            key=key,
            accept_file=True,
            file_type=["csv", "parquet", "xlsx"],
            accept_audio=True,
            audio_sample_rate=16000,
            max_upload_size=max(1, (max_bytes + 1_048_575) // 1_048_576),
            submit_mode="disable",
        )
        return review_voice_submission(submission, key)


def stage_attachment(client, submission, max_bytes):
    """Stage only the attached file; its question waits for schema confirmation."""
    uploaded = submission.files[0]
    if uploaded.size > max_bytes:
        raise APIError(f"File exceeds the {max_bytes:,}-byte upload limit.")
    if uploaded.name.lower().endswith(".xlsx"):
        content = uploaded.getvalue()
        info = client.inspect_workbook(content)
        st.session_state["pending_workbook"] = {
            "filename": uploaded.name,
            "content": content,
            "info": info,
            "question": submission.text,
        }
        return None
    return client.upload_file(uploaded.name, uploaded.getvalue())


def _column_type_editor(columns, *, key):
    return st.data_editor(
        [
            {
                "Column": c["name"],
                "Use as": TYPE_LABELS[c["suggested_type"]],
                "Missing values": c["null_count"],
                "Review note": c["warning"],
            }
            for c in columns
        ],
        hide_index=True,
        disabled=["Column", "Missing values", "Review note"],
        column_config={
            "Use as": st.column_config.SelectboxColumn(
                options=list(TYPE_LABELS.values()), required=True
            ),
            "Review note": st.column_config.TextColumn(width="large"),
        },
        key=key,
    )


def render_upload_review(client, upload):
    st.subheader("Confirm the data to analyze")
    st.caption(
        f"{upload['filename']} · {upload['row_count']:,} rows · {len(upload['columns'])} columns"
    )
    if upload.get("excel_selection"):
        selection = upload["excel_selection"]
        st.caption(
            f"Worksheet: {selection['sheet']} · table: {selection['cell_range']} (first row is the header)"
        )
    st.caption(
        "Check the preview and column notes, then confirm. Your question starts after this step."
    )
    if upload.get("warnings"):
        with st.expander("File notes", expanded=True):
            for warning in upload["warnings"]:
                st.write(warning)
    st.dataframe(
        upload["sample_rows"],
        hide_index=True,
        key=f"upload_preview_{upload['thread_id']}",
    )
    st.caption(
        "Preview: at most 10 rows. Type checks and declared key checks use the entire file."
    )
    notes = [c for c in upload["columns"] if c["warning"]]
    other = [c for c in upload["columns"] if not c["warning"]]
    names = [c["name"] for c in upload["columns"]]
    st.caption(
        f"Suggested types are ready for all {len(names)} columns. "
        "Keep IDs as text to preserve leading zeros. Choose the order of ambiguous dates explicitly. "
        "Types do not define business meaning or units."
    )
    with st.form(f"review_upload_{upload['thread_id']}"):
        selected = []
        if notes:
            st.markdown(f"**Columns with notes · {len(notes)}**")
            selected.extend(
                _column_type_editor(notes, key=f"upload_notes_{upload['thread_id']}")
            )
        if other:
            with st.expander(f"Other column types · {len(other)}", expanded=not notes):
                selected.extend(
                    _column_type_editor(
                        other, key=f"upload_types_{upload['thread_id']}"
                    )
                )
        with st.expander("Describe the rows (optional)"):
            grain = st.text_input(
                "What does one row represent?",
                max_chars=500,
                placeholder="For example, one order line; leave blank if unknown",
            )
            keys = st.multiselect(
                "Columns that uniquely identify a row",
                names,
                help="We check these identifiers against every row. Duplicates or missing identifiers must be corrected; rows are never silently removed.",
            )
        with st.expander("Original column details"):
            st.dataframe(
                [
                    {
                        "Column": c["name"],
                        "Stored type": c["stored_type"],
                        "Unique values": c["distinct_count"],
                    }
                    for c in upload["columns"]
                ],
                hide_index=True,
            )
        submitted = st.form_submit_button("Confirm data and continue", type="primary")
    if submitted:
        try:
            with st.spinner("Checking all rows and saving reviewed data…"):
                client.confirm_upload(
                    upload["thread_id"],
                    {
                        "types": {
                            row["Column"]: {
                                label: kind for kind, label in TYPE_LABELS.items()
                            }[row["Use as"]]
                            for row in selected
                        },
                        "grain": grain,
                        "key_columns": keys,
                    },
                )
        except APIError as exc:
            st.error(str(exc))
            return
        st.rerun()


def render_uploaded_source(client, upload):
    with st.expander("Your confirmed data"):
        st.write(
            f"{upload['filename']} · {upload['row_count']:,} rows · {len(upload['columns'])} columns"
        )
        if upload.get("excel_selection"):
            selection = upload["excel_selection"]
            st.caption(
                f"Worksheet: {selection['sheet']} · Selected table: {selection['cell_range']}"
            )
        review = upload["review"]
        st.write("One row represents: " + (review["grain"] or "Not specified"))
        st.write(
            "Row identifiers: " + (", ".join(review["key_columns"]) or "Not specified")
        )
        st.caption(
            "Analysis uses this saved file. When the source data was last updated is unknown; uploading it does not establish freshness."
        )
        st.dataframe(
            [
                {"Column": name, "Confirmed type": TYPE_LABELS[kind]}
                for name, kind in review["types"].items()
            ],
            hide_index=True,
        )
        for warning in upload.get("warnings", []):
            st.write(warning)
        st.link_button(
            "Download reviewed data", client.dataset_download_url(upload["result_id"])
        )
        with st.expander("Technical file details"):
            st.json(
                {
                    "sha256": upload["sha256"],
                    "review": review,
                    "excel_selection": upload.get("excel_selection"),
                    "import_warnings": upload.get("warnings", []),
                    "imported_at": upload["imported_at"],
                }
            )


def _table_scope(first_column, last_column, header_row, last_row):
    cell_range = f"{first_column.strip().upper()}{header_row}:{last_column.strip().upper()}{last_row}"
    try:
        left, top, right, bottom = range_boundaries(cell_range)
        if None in (left, top, right, bottom) or not 1 <= left <= right <= 16384:
            raise ValueError
    except ValueError as exc:
        raise ValueError(
            "Enter valid first and last column letters, such as A and D."
        ) from exc
    if not 1 <= top < bottom <= 1048576:
        raise ValueError(
            "The last data row must be below the row containing column names."
        )
    if right - left + 1 > MAX_UPLOAD_COLUMNS:
        raise ValueError(
            f"Choose at most {MAX_UPLOAD_COLUMNS} columns in Columns to include."
        )
    return cell_range, bottom - top, right - left + 1


def render_workbook_selection(client):
    pending = st.session_state["pending_workbook"]
    st.subheader("Choose the data in your workbook")
    st.caption(pending["filename"])
    st.caption("Step 1 of 2 · Choose one table. Next, check its data and column types.")
    if pending["question"]:
        st.info("Your question is saved and will start after you confirm the data.")
    sheets = {s["name"]: s for s in pending["info"]["sheets"]}
    sheet = st.selectbox("Worksheet with your data", list(sheets))
    preview = sheets[sheet]["preview"]
    st.dataframe(
        pd.DataFrame(
            preview,
            index=pd.Index(range(1, len(preview) + 1), name="Excel row"),
            columns=[
                get_column_letter(i + 1)
                for i in range(max((len(r) for r in preview), default=0))
            ],
        ),
    )
    st.caption(
        "Preview: first 10 worksheet rows and 20 columns. Use the row numbers to skip titles and exclude totals or footers. All selected rows are retained."
    )
    left, top, right, bottom = range_boundaries(sheets[sheet]["range"])
    key = hashlib.sha256(pending["content"]).hexdigest()[:12] + "_" + sheet
    with st.container(horizontal=True):
        header_row = st.number_input(
            "Column names are on row",
            min_value=1,
            max_value=1048575,
            value=top,
            key=f"excel_header_{key}",
        )
        last_row = st.number_input(
            "Last data row",
            min_value=1,
            max_value=1048576,
            value=bottom,
            key=f"excel_last_{key}",
        )
    with st.expander("Columns to include"):
        with st.container(horizontal=True):
            first_column = st.text_input(
                "First column",
                value=get_column_letter(left),
                key=f"excel_first_column_{key}",
            )
            last_column = st.text_input(
                "Last column",
                value=get_column_letter(right),
                key=f"excel_last_column_{key}",
            )
        st.caption("Use the worksheet letters shown above, for example A through D.")
    try:
        cell_range, row_count, column_count = _table_scope(
            first_column, last_column, header_row, last_row
        )
        st.success(f"Selected data: {row_count:,} rows · {column_count} columns")
        st.caption(
            f"Column names: row {header_row} · Worksheet range: {sheet}!{cell_range}"
        )
        valid = True
    except ValueError as exc:
        st.error(str(exc))
        valid = False
    with st.form(f"excel_table_selection_{key}"):
        with st.expander("If your table contains formulas"):
            cached = st.checkbox(
                "Use values saved by Excel for formula cells",
                help="The app does not recalculate formulas. Saved values may be out of date. Recalculate and save the workbook in Excel first if needed.",
            )
        submitted = st.form_submit_button(
            "Continue to data review", type="primary", disabled=not valid
        )
    if submitted:
        try:
            with st.spinner("Loading the selected rows for review…"):
                upload = client.upload_file(
                    pending["filename"],
                    pending["content"],
                    {
                        "sheet": sheet,
                        "cell_range": cell_range,
                        "use_cached_formulas": cached,
                    },
                )
        except APIError as exc:
            st.error(str(exc))
        else:
            st.session_state.setdefault("upload_questions", {})[upload["thread_id"]] = (
                pending["question"]
            )
            del st.session_state["pending_workbook"]
            st.query_params["thread_id"] = upload["thread_id"]
            st.session_state["current_thread_id"] = None
            st.rerun()
    if st.button("Cancel workbook import"):
        del st.session_state["pending_workbook"]
        st.rerun()
